"""Lettura e scrittura atomica di finanze.xlsx, senza altri database.

Un lock serializza le sessioni Streamlit. Ogni transazione rilegge il file,
modifica tutti i fogli coinvolti e sostituisce l'Excel in un'unica operazione.
"""
from __future__ import annotations

import os
import shutil
import tempfile
from contextlib import contextmanager
from pathlib import Path
from zipfile import BadZipFile

from filelock import FileLock, Timeout
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .errors import ConflictError, StorageError
from .utils import new_id


SCHEMA = {
    "Movimenti": [
        "ID", "Data", "Tipo", "Importo", "Categoria", "Descrizione", "Fonte",
        "Note", "Data_creazione", "Conto", "Arbitraggio_ID", "Conto_destinazione",
    ],
    "Arbitraggio": [
        "ID", "Data partita", "Numero pacco", "Squadra casa", "Squadra ospite",
        "Compenso", "Note", "Conto", "Movimento_ID", "Categoria", "Stato", "Data incasso", "Km", "Categoria partita",
    ],
    "Categorie": ["ID", "Nome", "Tipo"],
    "Conti": ["ID", "Nome", "Saldo_iniziale"],
}

DEFAULT_CATEGORIES = [
    ("Cibo", "Uscita"), ("Macchina", "Uscita"), ("Benzina", "Uscita"),
    ("Vestiti", "Uscita"), ("Svago", "Uscita"), ("Stipendio", "Entrata"),
    ("Arbitraggio", "Entrata"), ("Altro", "Entrambe"),
]


class ExcelStore:
    def __init__(self, path: str | Path):
        self.path = Path(path).expanduser().resolve()
        self.lock = FileLock(str(self.path) + ".lock", timeout=10)
        self.backup_path = self.path.with_name(f"{self.path.stem}.backup.xlsx")

    @contextmanager
    def _locked(self):
        try:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            with self.lock:
                yield
        except Timeout as exc:
            raise StorageError("Il file è in uso dall'app. Riprova tra qualche secondo.") from exc
        except PermissionError as exc:
            raise StorageError("Non posso salvare l'Excel. Chiudi il file in Excel e riprova.") from exc
        except OSError as exc:
            raise StorageError("Non posso leggere o salvare l'archivio. Controlla la cartella e lo spazio disponibile.") from exc

    def _fingerprint(self):
        if not self.path.exists():
            return None
        stat = self.path.stat()
        return stat.st_mtime_ns, stat.st_size

    def _load(self):
        fingerprint = self._fingerprint()
        if fingerprint is None:
            workbook = Workbook()
            workbook.remove(workbook.active)
        else:
            try:
                workbook = load_workbook(self.path)
            except (OSError, PermissionError):
                raise
            except Exception as exc:
                raise StorageError("Il file Excel non è leggibile. Non è stato sovrascritto; puoi recuperare finanze.backup.xlsx.") from exc
        changed = fingerprint is None
        for name, columns in SCHEMA.items():
            if name not in workbook.sheetnames:
                worksheet = workbook.create_sheet(name)
                worksheet.append(columns)
                self._seed(name, worksheet)
                changed = True
            else:
                worksheet = workbook[name]
                headers = [cell.value for cell in worksheet[1]]
                if not any(headers):
                    for index, column in enumerate(columns, 1):
                        worksheet.cell(1, index, column)
                    changed = True
                else:
                    for column in columns:
                        if column not in headers:
                            worksheet.cell(1, worksheet.max_column + 1, column)
                            changed = True
            if changed:
                self._style(worksheet)
        return workbook, fingerprint, changed

    @staticmethod
    def _seed(name, worksheet):
        if name == "Categorie":
            for label, kind in DEFAULT_CATEGORIES:
                worksheet.append([new_id("C"), label, kind])
        elif name == "Conti":
            for label in ("Conto personale", "Contanti"):
                worksheet.append([new_id("A"), label, 0.0])

    @staticmethod
    def _rows(workbook):
        tables = {}
        for name in SCHEMA:
            worksheet = workbook[name]
            headers = [cell.value for cell in worksheet[1]]
            populated = [header for header in headers if header is not None]
            if len(populated) != len(set(populated)):
                raise StorageError(f"Nel foglio {name} ci sono intestazioni duplicate. Correggile in Excel.")
            records = []
            for values in worksheet.iter_rows(min_row=2, values_only=True):
                if any(value is not None for value in values):
                    records.append({header: value for header, value in zip(headers, values) if header is not None})
            tables[name] = records
        return tables

    @staticmethod
    def _style(worksheet):
        headers = [cell.value for cell in worksheet[1]]
        worksheet.freeze_panes = "B2"
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.sheet_properties.tabColor = "14243B"
        worksheet.row_dimensions[1].height = 30
        worksheet.auto_filter.ref = f"A1:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
        for cell in worksheet[1]:
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="14243B")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for index, header in enumerate(headers, 1):
            width = 22
            if header in ("ID", "Movimento_ID", "Arbitraggio_ID"):
                width = 37
            elif header in ("Descrizione", "Note"):
                width = 46
            elif header in ("Importo", "Compenso", "Saldo_iniziale"):
                width = 20
            elif header in ("Data", "Data partita", "Data incasso"):
                width = 18
            elif header == "Data_creazione":
                width = 23
            worksheet.column_dimensions[get_column_letter(index)].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell, header in zip(row, headers):
                cell.font = Font(name="Arial", size=10, color="14243B")
                cell.alignment = Alignment(vertical="center")
                cell.fill = PatternFill("solid", fgColor="F2F6FA" if cell.row % 2 == 0 else "FFFFFF")
                if header in ("Importo", "Compenso", "Saldo_iniziale"):
                    cell.number_format = '#,##0.00" €"'
                elif header in ("Data", "Data partita", "Data incasso"):
                    cell.number_format = "dd/mm/yyyy"
                elif header == "Data_creazione":
                    cell.number_format = "dd/mm/yyyy hh:mm:ss"
                elif header == "Km":
                    cell.number_format = "0.0"

    def _write_rows(self, workbook, tables):
        for name, records in tables.items():
            worksheet = workbook[name]
            headers = [cell.value for cell in worksheet[1]]
            if worksheet.max_row > 1:
                worksheet.delete_rows(2, worksheet.max_row - 1)
            for record in records:
                worksheet.append([record.get(header) for header in headers])
            # Treat user-provided text literally, including strings beginning with '='.
            for row in worksheet.iter_rows(min_row=2):
                for cell, header in zip(row, headers):
                    if header in SCHEMA[name] and isinstance(cell.value, str):
                        cell.data_type = "s"
            for table in worksheet.tables.values():
                table.ref = f"A1:{get_column_letter(worksheet.max_column)}{max(2, worksheet.max_row)}"
            self._style(worksheet)

    def _save(self, workbook, fingerprint):
        descriptor, temporary_name = tempfile.mkstemp(prefix=".finanze-", suffix=".xlsx", dir=self.path.parent)
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            workbook.save(temporary_path)
            with temporary_path.open("rb") as handle:
                os.fsync(handle.fileno())
            if self._fingerprint() != fingerprint:
                raise ConflictError("L'Excel è cambiato durante il salvataggio. Aggiorna i dati e riprova.")
            if fingerprint is not None:
                backup_tmp = temporary_path.with_suffix(".backup")
                try:
                    shutil.copy2(self.path, backup_tmp)
                    os.replace(backup_tmp, self.backup_path)
                finally:
                    backup_tmp.unlink(missing_ok=True)
            os.replace(temporary_path, self.path)
        finally:
            temporary_path.unlink(missing_ok=True)

    def read(self):
        with self._locked():
            workbook, fingerprint, changed = self._load()
            try:
                records = self._rows(workbook)
                if changed:
                    self._save(workbook, fingerprint)
                return records
            finally:
                workbook.close()

    def transaction(self, operation):
        with self._locked():
            workbook, fingerprint, _ = self._load()
            try:
                tables = self._rows(workbook)
                result = operation(tables)
                self._write_rows(workbook, tables)
                self._save(workbook, fingerprint)
                return result
            finally:
                workbook.close()

    def download(self) -> bytes:
        with self._locked():
            return self.path.read_bytes()
