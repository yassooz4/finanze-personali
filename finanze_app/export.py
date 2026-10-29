"""Excel exports are copies, never the production database."""
from io import BytesIO
import pandas as pd
from openpyxl import Workbook
from .storage import SCHEMA, ExcelStore


def _movements_export_frame(frame):
    """Copia pulita dei movimenti: niente colonne interne e protezione dalle formule."""
    data = frame[[column for column in frame.columns if not str(column).startswith("_")]].copy()
    for column in data.columns:
        if str(data[column].dtype).startswith("datetime"):
            data[column] = data[column].dt.strftime("%Y-%m-%d")
        elif pd.api.types.is_object_dtype(data[column]) or pd.api.types.is_string_dtype(data[column]):
            data[column] = data[column].map(
                lambda value: "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value
            )
    return data


def movements_csv(frame):
    """CSV dei movimenti (separatore ';' e UTF-8 con BOM, apribile direttamente in Excel)."""
    return _movements_export_frame(frame).to_csv(index=False, sep=";", decimal=",").encode("utf-8-sig")


def movements_xlsx(frame):
    """File Excel con i soli movimenti passati."""
    output = BytesIO()
    _movements_export_frame(frame).to_excel(output, index=False, sheet_name="Movimenti", engine="openpyxl")
    return output.getvalue()


def excel_export(tables):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, schema in SCHEMA.items():
        sheet = workbook.create_sheet(name)
        headers = list(schema)
        for record in tables[name]:
            headers.extend(key for key in record if key not in headers)
        sheet.append(headers)
        for record in tables[name]:
            sheet.append([record.get(key) for key in headers])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
        ExcelStore._style(sheet)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
