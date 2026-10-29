"""Report completo scaricabile (Excel/CSV): copia filtrata dei dati, mai il database di produzione."""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analytics as stats
from .export import movements_csv
from .service import Snapshot
from .utils import cents, timestamp

SECTIONS = ["Riepilogo", "Movimenti", "Categorie", "Mesi", "Conti", "Arbitraggio"]
KINDS = ("Entrata", "Uscita", "Giroconto")

MONEY_COLUMNS = {"Importo", "Saldo progressivo", "Entrate", "Uscite", "Differenza", "Saldo iniziale periodo", "Saldo finale periodo", "Compenso"}
PERCENT_COLUMNS = {"% sul totale"}
DATE_COLUMNS = {"Data", "Data partita", "Data incasso"}
MONEY_FORMAT = '#,##0.00 "€"'
MATCH_COLUMNS = ["Data partita", "Numero pacco", "Squadra casa", "Squadra ospite", "Compenso", "Stato", "Data incasso", "Conto", "Categoria", "Km", "Categoria partita", "Note"]


@dataclass
class Report:
    summary: list  # (voce, valore, formato) con formato in {"text", "money", "int"}
    movements: pd.DataFrame
    categories: pd.DataFrame
    months: pd.DataFrame
    accounts: pd.DataFrame
    matches: pd.DataFrame

    def frame(self, section: str) -> pd.DataFrame:
        return {"Movimenti": self.movements, "Categorie": self.categories, "Mesi": self.months,
                "Conti": self.accounts, "Arbitraggio": self.matches}[section]


def _opening_cents(snapshot: Snapshot, account=None) -> int:
    return sum(cents(row["Saldo_iniziale"]) for row in snapshot.accounts if account is None or row["Nome"] == account)


def _balance_cents(snapshot: Snapshot, day, account=None) -> int:
    """Saldo a fine giornata `day`; con `day=None` è il saldo più recente. Prima del primo movimento vale il saldo iniziale."""
    history = stats.balance_history(snapshot, end=day, account=account)
    return cents(history["Saldo"].iloc[-1]) if not history.empty else _opening_cents(snapshot, account)


def _ledger(snapshot: Snapshot, account=None) -> pd.DataFrame:
    """Tutti i movimenti dell'ambito scelto in ordine cronologico, con saldo progressivo sull'intero storico."""
    view = stats.movement_view(snapshot.movements, account)
    view = view.sort_values(["Data", "Data_creazione", "ID"]).reset_index(drop=True)
    sign = view["Tipo"].map({"Entrata": 1, "Uscita": -1}).fillna(0).astype("int64")
    delta = view["Importo"].map(cents).astype("int64") * sign
    view["Saldo progressivo"] = (delta.cumsum() + _opening_cents(snapshot, account)) / 100
    return view


def _movements_table(view: pd.DataFrame, account) -> pd.DataFrame:
    columns = ["Data", "Tipo", "Importo", "Categoria", "Descrizione", "Conto", "Conto destinazione", "Fonte", "Note", "Origine", "Saldo progressivo"]
    if view.empty:
        return pd.DataFrame(columns=columns)
    out = pd.DataFrame(index=view.index)
    transfer = view["_Giroconto"]
    tipo = view["Tipo"]
    if account is not None:
        suffix = view["Tipo"].map({"Entrata": "in entrata", "Uscita": "in uscita"}).fillna("")
        tipo = tipo.where(~transfer, "Giroconto " + suffix)
    out["Data"] = view["Data"]
    out["Tipo"] = tipo
    out["Importo"] = (view["Importo"] * np.where(view["Tipo"] == "Uscita", -1, 1)).round(2)
    out["Categoria"] = view["Categoria"].fillna("")
    out["Descrizione"] = view["Descrizione"].fillna("")
    out["Conto"] = view["Conto"].fillna("")
    out["Conto destinazione"] = view["Conto_destinazione"].fillna("")
    out["Fonte"] = view["Fonte"].fillna("")
    out["Note"] = view["Note"].fillna("")
    out["Origine"] = np.where(view["Arbitraggio_ID"].fillna("").astype(bool), "Partita", np.where(transfer, "Giroconto", "Manuale"))
    out["Saldo progressivo"] = view["Saldo progressivo"]
    return out[columns].reset_index(drop=True)


def _categories_table(view: pd.DataFrame) -> pd.DataFrame:
    columns = ["Tipo", "Categoria", "Importo", "% sul totale", "Movimenti"]
    work = view[~view["_Giroconto"]].copy()
    if work.empty:
        return pd.DataFrame(columns=columns)
    work["Categoria"] = work["Categoria"].replace("", "Senza categoria").fillna("Senza categoria")
    work["Centesimi"] = work["Importo"].map(cents)
    grouped = work.groupby(["Tipo", "Categoria"], as_index=False).agg(Centesimi=("Centesimi", "sum"), Movimenti=("ID", "count"))
    total = grouped.groupby("Tipo")["Centesimi"].transform("sum")
    grouped["Importo"] = grouped["Centesimi"] / 100
    grouped["% sul totale"] = (grouped["Centesimi"] / total).where(total > 0, 0.0)
    grouped["_ordine"] = grouped["Tipo"].map({"Entrata": 0, "Uscita": 1}).fillna(2)
    grouped = grouped.sort_values(["_ordine", "Importo"], ascending=[True, False])
    return grouped[columns].reset_index(drop=True)


def _months_table(view: pd.DataFrame, start, end) -> pd.DataFrame:
    flows = stats.monthly_flows(view, start, end)
    if flows.empty:
        return pd.DataFrame(columns=["Mese", "Entrate", "Uscite", "Differenza"])
    flows = flows.copy()
    flows["Mese"] = [f"{stats.MONTHS[month.month - 1]} {month.year}" for month in flows["Mese"]]
    return flows[["Mese", "Entrate", "Uscite", "Differenza"]]


def _accounts_table(snapshot: Snapshot, start, end, account) -> pd.DataFrame:
    rows = []
    before = pd.Timestamp(start) - pd.Timedelta(days=1) if start is not None else None
    for name in [account] if account else snapshot.account_names():
        view = stats.date_filter(stats.movement_view(snapshot.movements, name), start, end)
        income, expenses, _ = stats.totals(view)
        opening = _balance_cents(snapshot, before, name) if before is not None else _opening_cents(snapshot, name)
        closing = _balance_cents(snapshot, end, name)
        rows.append({"Conto": name, "Saldo iniziale periodo": opening / 100, "Entrate": income, "Uscite": expenses, "Saldo finale periodo": closing / 100})
    return pd.DataFrame(rows, columns=["Conto", "Saldo iniziale periodo", "Entrate", "Uscite", "Saldo finale periodo"])


def _matches_table(snapshot: Snapshot, start, end, account) -> pd.DataFrame:
    matches = stats.date_filter(snapshot.matches, start, end, column="Data partita")
    if account:
        matches = matches[matches["Conto"] == account]
    return matches.sort_values("Data partita")[MATCH_COLUMNS].reset_index(drop=True)


def _day(value) -> str:
    return value.strftime("%d/%m/%Y")


def period_label(start, end) -> str:
    if start is None and end is None:
        return "Tutto lo storico"
    return f"dal {_day(start) if start is not None else 'primo movimento'} al {_day(end) if end is not None else 'oggi'}"


def file_stem(start, end) -> str:
    if start is None and end is None:
        return "report_finanze_storico"
    return f"report_finanze_{start or 'inizio'}_{end or 'oggi'}"


def build_report(snapshot: Snapshot, *, start=None, end=None, account=None, kind=None, category=None) -> Report:
    """`category=""` seleziona i movimenti senza categoria; `None` non filtra.

    Tipo e categoria filtrano Movimenti, Categorie, Mesi e i totali del Riepilogo.
    Saldi, Conti e Arbitraggio dipendono solo da periodo e conto.
    """
    view = stats.date_filter(_ledger(snapshot, account), start, end)
    if kind == "Giroconto":
        view = view[view["_Giroconto"]]
    elif kind is not None:
        view = view[view["Tipo"] == kind]
    if category is not None:
        view = view[view["Categoria"].fillna("") == category]
    income, expenses, net = stats.totals(view)
    matches = _matches_table(snapshot, start, end, account)
    before = pd.Timestamp(start) - pd.Timedelta(days=1) if start is not None else None
    opening = _balance_cents(snapshot, before, account) if before is not None else _opening_cents(snapshot, account)
    closing = _balance_cents(snapshot, end, account)
    summary = [
        ("Generato il", timestamp().strftime("%d/%m/%Y %H:%M"), "text"),
        ("Periodo", period_label(start, end), "text"),
        ("Conto", account or "Tutti i conti", "text"),
        ("Tipo", kind or "Tutti", "text"),
        ("Categoria", "Tutte" if category is None else (category or "Senza categoria"), "text"),
        ("Movimenti nel report", len(view), "int"),
        ("Entrate", income, "money"),
        ("Uscite", expenses, "money"),
        ("Differenza (entrate − uscite)", net, "money"),
        ("Saldo a inizio periodo", opening / 100, "money"),
        ("Saldo a fine periodo", closing / 100, "money"),
        ("Saldo attuale", _balance_cents(snapshot, None, account) / 100, "money"),
        ("Partite arbitraggio nel periodo", len(matches), "int"),
        ("Compensi arbitraggio nel periodo", sum(cents(value) for value in matches["Compenso"]) / 100, "money"),
        ("Nota", "Con un conto selezionato i giroconti contano come entrate/uscite di quel conto e non compaiono per categoria. "
                 "Saldi, Conti e Arbitraggio non dipendono dai filtri tipo e categoria.", "text"),
    ]
    return Report(summary, _movements_table(view, account), _categories_table(view),
                  _months_table(view, start, end), _accounts_table(snapshot, start, end, account), matches)


def _cell_value(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, np.generic):
        return value.item()
    return value


def _put(ws, row, column, value):
    cell = ws.cell(row=row, column=column, value=_cell_value(value))
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"  # testo, mai formula
    return cell


def _style_sheet(ws, header_row=1):
    fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None and not (isinstance(c.value, str) and len(c.value) > 80)), default=8)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 2, 10), 60)


def _write_frame(ws, frame: pd.DataFrame):
    ws.append(list(frame.columns))
    for r, record in enumerate(frame.itertuples(index=False), start=2):
        for c, value in enumerate(record, start=1):
            cell = _put(ws, r, c, value)
            name = frame.columns[c - 1]
            if name in MONEY_COLUMNS:
                cell.number_format = MONEY_FORMAT
            elif name in PERCENT_COLUMNS:
                cell.number_format = "0.0%"
            elif name in DATE_COLUMNS and cell.value is not None:
                cell.number_format = "DD/MM/YYYY"
    _style_sheet(ws)


def _write_summary(ws, summary):
    ws.append(["Voce", "Valore"])
    for r, (label, value, kind) in enumerate(summary, start=2):
        _put(ws, r, 1, label).font = Font(bold=True)
        cell = _put(ws, r, 2, value)
        cell.alignment = Alignment(horizontal="left", wrap_text=kind == "text" and len(str(value)) > 60)
        if kind == "money":
            cell.number_format = MONEY_FORMAT
    _style_sheet(ws)
    ws.column_dimensions["B"].width = 60


def to_excel(report: Report, sections=SECTIONS) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in SECTIONS:
        if name not in sections:
            continue
        sheet = workbook.create_sheet(name)
        if name == "Riepilogo":
            _write_summary(sheet, report.summary)
        else:
            _write_frame(sheet, report.frame(name))
    if not workbook.sheetnames:
        workbook.create_sheet("Riepilogo")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def to_csv(report: Report) -> bytes:
    return movements_csv(report.movements)
