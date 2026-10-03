from copy import deepcopy
from io import BytesIO
from datetime import date, timedelta

import pytest
from requests.exceptions import Timeout, ConnectionError
from openpyxl import load_workbook

from finanze_app.sheets_storage import GoogleSheetsStore
from finanze_app.service import FinanceService
from finanze_app.errors import StorageError, ConflictError, ValidationError
from finanze_app.analytics import total_balance
from finanze_app.export import excel_export
from finanze_app.storage import SCHEMA
from finanze_app.utils import today


class RemoteSheet:
    """API emulator: state exists in the remote object, not the adapter."""
    id = 'test_spreadsheet_with_long_id'

    def __init__(self):
        self.meta = {'Privato': {'sheetId': 0, 'title': 'Privato', 'gridProperties': {'rowCount': 1000, 'columnCount': 26}}}
        self.grids = {'Privato': [['conservare']]}
        self.fail = None
        self.after_read = None
        self.calls = []

    def fetch_sheet_metadata(self, params=None):
        return {'sheets': [{'properties': deepcopy(m)} for m in self.meta.values()]}

    def values_batch_get(self, ranges, params=None):
        result = {'valueRanges': [{'values': deepcopy(self.grids[r.strip("'")])} for r in ranges]}
        if self.after_read:
            callback, self.after_read = self.after_read, None
            callback()
        return result

    def batch_update(self, body):
        if self.fail == 'before':
            self.fail = None
            raise ConnectionError('secret must never be displayed')
        meta, grids = deepcopy(self.meta), deepcopy(self.grids)
        for req in body['requests']:
            if 'addSheet' in req:
                m = req['addSheet']['properties']
                meta[m['title']] = deepcopy(m)
                grids[m['title']] = []
            elif 'appendDimension' in req:
                item = req['appendDimension']
                m = next(m for m in meta.values() if m['sheetId'] == item['sheetId'])
                m['gridProperties']['rowCount' if item['dimension'] == 'ROWS' else 'columnCount'] += item['length']
            elif 'updateCells' in req:
                item = req['updateCells']
                span = item.get('range', item.get('start'))
                name = next(n for n,m in meta.items() if m['sheetId'] == span['sheetId'])
                start = span.get('startRowIndex', span.get('rowIndex', 0))
                col = span.get('startColumnIndex', span.get('columnIndex', 0))
                rows = item.get('rows', [])
                end = span.get('endRowIndex', start+len(rows))
                width = span.get('endColumnIndex', col+max([len(r['values']) for r in rows]+[0]))
                grid = grids[name]
                while len(grid) < end:
                    grid.append([])
                for index in range(start, end):
                    while len(grid[index]) < width:
                        grid[index].append('')
                    cells = rows[index-start]['values'] if index-start < len(rows) else []
                    for c in range(col, width):
                        value = cells[c-col].get('userEnteredValue', {}) if c-col < len(cells) else {}
                        grid[index][c] = next(iter(value.values()), '')
                while grid and not any(v != '' for v in grid[-1]):
                    grid.pop()
        self.meta, self.grids = meta, grids
        self.calls.append(deepcopy(body))
        if self.fail == 'after':
            self.fail = None
            raise Timeout('accepted remotely')
        return {}


def service(remote):
    return FinanceService(None, store=GoogleSheetsStore(remote))


def transfer(s, **changes):
    args = dict(kind='Giroconto', amount=30, date=today(), description='', category='', source='', account='Conto personale', destination_account='Contanti')
    args.update(changes)
    return s.save_movement(**args)


def test_existing_sheet_migrates_without_losing_movements_matches_or_custom_data():
    remote = RemoteSheet()
    s = service(remote)
    match(s, status='Ricevuto', received_date=date(2026,9,30), account='Contanti', category='Arbitraggio')
    before = s.snapshot()
    headers = remote.grids['Movimenti'][0]
    index = headers.index('Conto_destinazione')
    for row in remote.grids['Movimenti']:
        if len(row) > index:
            row.pop(index)
    migrated = service(remote).snapshot()
    assert migrated.tables == before.tables
    assert 'Conto_destinazione' in remote.grids['Movimenti'][0]
    assert total_balance(migrated) == 45
    assert remote.grids['Privato'] == [['conservare']]
    mid = transfer(service(remote), account='Contanti', destination_account='Conto personale', amount=20)
    reloaded = service(remote).snapshot()
    assert len(reloaded.movements) == 2 and len(reloaded.matches) == 1
    assert mid in reloaded.movements['ID'].tolist()
    assert total_balance(reloaded) == 45


def test_transfer_restart_edit_rename_and_delete_preserve_balances():
    from finanze_app.analytics import account_balances, totals
    remote = RemoteSheet()
    s = service(remote)
    source, target = s.snapshot().accounts
    s.save_account(name=source['Nome'], opening_balance=100, account_id=source['ID'], expected=source)
    mid = transfer(s)
    row = service(remote).snapshot().tables['Movimenti'][0]
    transfer(service(remote), movement_id=mid, expected=row, amount=40)
    restarted = service(remote)
    current = restarted.snapshot()
    target = next(a for a in current.accounts if a['ID'] == target['ID'])
    restarted.save_account(name='Portafoglio', opening_balance=target['Saldo_iniziale'], account_id=target['ID'], expected=target)
    current = service(remote).snapshot()
    record = current.tables['Movimenti'][0]
    assert record['ID'] == mid and record['Conto_destinazione'] == 'Portafoglio'
    assert len(current.movements) == 1 and total_balance(current) == 100
    assert account_balances(current).set_index('Conto')['Saldo'].to_dict() == {'Conto personale': 60, 'Portafoglio': 40}
    assert totals(current.movements) == (0, 0, 0)
    service(remote).delete_movement(mid, expected=record)
    current = service(remote).snapshot()
    assert current.movements.empty and total_balance(current) == 100
    assert account_balances(current).set_index('Conto')['Saldo'].to_dict() == {'Conto personale': 100, 'Portafoglio': 0}


@pytest.mark.parametrize('failure', ['before', 'after'])
def test_transfer_timeout_preserves_atomicity(failure):
    remote = RemoteSheet()
    s = service(remote)
    s.snapshot()
    remote.fail = failure
    if failure == 'before':
        with pytest.raises(StorageError, match='non confermato'):
            transfer(s)
        assert service(remote).snapshot().movements.empty
    else:
        mid = transfer(s)
        current = service(remote).snapshot()
        assert len(current.movements) == 1 and current.tables['Movimenti'][0]['ID'] == mid
    assert total_balance(service(remote).snapshot()) == 0


def match(s, **changes):
    args = dict(date=date(2026,9,30), package='0012', home='Casa', away='Ospiti', fee=45, notes='', km=27.5, match_category='Under 17')
    args.update(changes)
    return s.save_match(**args)


def test_new_archive_restart_and_payment_roundtrip():
    remote = RemoteSheet()
    s = service(remote)
    assert s.snapshot().matches.empty
    mid = match(s)
    restarted = service(remote)
    snap = restarted.snapshot()
    row = snap.tables['Arbitraggio'][0]
    assert row['ID'] == mid and row['Numero pacco'] == '0012'
    assert row['Km'] == 27.5 and row['Categoria partita'] == 'Under 17'
    assert row['Stato'] == 'Da ricevere' and total_balance(snap) == 0
    match(restarted, match_id=mid, status='Ricevuto', received_date=date(2026,9,30), account='Contanti', category='Arbitraggio', expected=row)
    snap = service(remote).snapshot()
    assert total_balance(snap) == 45 and len(snap.movements) == 1
    assert snap.tables['Movimenti'][0]['Data'] == date(2026,9,30)
    last = remote.calls[-1]['requests']
    touched = {r['updateCells']['range']['sheetId'] for r in last if 'updateCells' in r}
    assert touched == {remote.meta[n]['sheetId'] for n in ['Movimenti','Arbitraggio']}
    match(s, match_id=mid, status='Ricevuto', received_date=date(2026,9,30), account='Contanti', category='Arbitraggio', fee=55)
    assert total_balance(s.snapshot()) == 55 and len(s.snapshot().movements) == 1
    match(s, match_id=mid, status='Da ricevere')
    assert s.snapshot().movements.empty and total_balance(s.snapshot()) == 0
    s.delete_match(mid)
    assert service(remote).snapshot().matches.empty
    assert remote.grids['Privato'] == [['conservare']]


def test_future_match_persists_without_credit_and_is_paid_after_match(monkeypatch):
    remote = RemoteSheet()
    tomorrow = today() + timedelta(days=1)
    mid = match(service(remote), date=tomorrow)
    restarted = service(remote)
    snap = restarted.snapshot()
    row = snap.tables['Arbitraggio'][0]
    assert row['ID'] == mid and row['Data partita'] == tomorrow
    assert row['Stato'] == 'Da ricevere'
    assert snap.movements.empty and total_balance(snap) == 0
    later = tomorrow + timedelta(days=7)
    match(restarted, date=later, match_id=mid, expected=row)
    row = service(remote).snapshot().tables['Arbitraggio'][0]
    assert row['Data partita'] == later
    before = deepcopy(remote.grids)
    for payment_date in (today(), later):
        with pytest.raises(ValidationError):
            service(remote).update_match_payments([{
                'ID': mid, 'Stato': 'Ricevuto', 'Data incasso': payment_date,
                'Conto': 'Contanti', 'expected': row,
            }])
        assert remote.grids == before
    monkeypatch.setattr('finanze_app.utils.today', lambda: later)
    service(remote).update_match_payments([{
        'ID': mid, 'Stato': 'Ricevuto', 'Data incasso': later,
        'Conto': 'Contanti', 'expected': row,
    }])
    snap = service(remote).snapshot()
    assert len(snap.matches) == 1 and len(snap.movements) == 1
    assert snap.tables['Arbitraggio'][0]['ID'] == mid
    assert snap.tables['Arbitraggio'][0]['Stato'] == 'Ricevuto'
    assert snap.tables['Movimenti'][0]['Data'] == later
    assert total_balance(snap) == 45


@pytest.mark.parametrize('failure', ['before','after'])
def test_timeout_does_not_duplicate_or_report_local_success(failure):
    remote = RemoteSheet(); s = service(remote); s.snapshot()
    remote.fail = failure
    if failure == 'before':
        with pytest.raises(StorageError, match='non confermato'):
            match(s)
        assert service(remote).snapshot().matches.empty
    else:
        match(s)
        assert len(service(remote).snapshot().matches) == 1


def test_external_change_aborts_save():
    remote = RemoteSheet(); s = service(remote); s.snapshot()
    remote.after_read = lambda: remote.grids['Conti'][1].__setitem__(1, 'Cambiato')
    with pytest.raises(ConflictError):
        match(s)
    assert service(remote).snapshot().matches.empty


def test_literal_formula_export_and_renames():
    remote = RemoteSheet(); s = service(remote); snap = s.snapshot()
    s.save_movement(kind='Uscita', amount=12, date=date(2026,9,30), description='=1+1', category='Cibo', source='Test', account='Conto personale')
    cat = next(c for c in snap.categories if c['Nome'] == 'Cibo')
    s.save_category(name='Alimentari', kind='Uscita', category_id=cat['ID'], expected=cat)
    assert service(remote).snapshot().tables['Movimenti'][0]['Categoria'] == 'Alimentari'
    data = excel_export(s.snapshot().tables)
    book = load_workbook(BytesIO(data))
    assert book['Movimenti']['F2'].value == '=1+1'
    assert book['Movimenti']['F2'].data_type == 's'
    assert book['Movimenti']['D2'].value == 12
    book.close()


def test_malformed_header_no_overwrite():
    remote = RemoteSheet(); s = service(remote); s.snapshot()
    original = deepcopy(remote.grids)
    remote.grids['Movimenti'][0][1] = 'ID'
    with pytest.raises(StorageError, match='duplicate'):
        match(s)
    assert remote.grids['Arbitraggio'] == original['Arbitraggio']


def test_extra_columns_and_growth_preserved():
    remote = RemoteSheet(); s = service(remote); s.snapshot()
    mid = match(s)
    remote.grids['Arbitraggio'][0].append('Personale')
    remote.grids['Arbitraggio'][1].append('=1+1')
    remote.meta['Arbitraggio']['gridProperties']['rowCount'] = 2
    fresh = service(remote)
    match(fresh)
    assert remote.meta['Arbitraggio']['gridProperties']['rowCount'] == 3
    match(fresh, match_id=mid, fee=99)
    assert remote.grids['Arbitraggio'][1][-1] == '=1+1'
    assert remote.meta['Arbitraggio']['gridProperties']['rowCount'] == 3


def test_read_error_never_returns_empty_archive():
    remote = RemoteSheet(); s = service(remote); s.snapshot(); match(s)
    remote.values_batch_get = lambda *a, **kw: (_ for _ in ()).throw(ConnectionError('private detail'))
    with pytest.raises(StorageError) as exc:
        s.snapshot()
    assert 'private detail' not in str(exc.value)
    assert len(remote.grids['Arbitraggio']) == 2


def test_streamlit_pages_use_remote_sheet(monkeypatch):
    from pathlib import Path
    from streamlit.testing.v1 import AppTest
    from finanze_app import sheets_storage
    remote = RemoteSheet()
    monkeypatch.delenv('FINANZE_FILE', raising=False)
    monkeypatch.setattr(sheets_storage, 'google_store', lambda *args: GoogleSheetsStore(remote))
    app = AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'), default_timeout=30)
    app.secrets['accesso'] = {'password': 'test-password'}
    app.secrets['sheets'] = {'spreadsheet_id': remote.id}
    app.secrets['gcp_service_account'] = {'client_email': 'test'}
    app.run()
    app.text_input(key='_login_password').set_value('test-password')
    app.button[0].click().run()
    assert not app.exception and not app.error
    service(remote).save_match(date=date(2026,9,30), package='007', home='A', away='B', fee=40, km=10, match_category='Juniores')
    for page in ['⚽ Arbitraggio','💸 Movimenti','📊 Statistiche','⚙️ Gestione','🏠 Dashboard']:
        app.sidebar.radio[0].set_value(page).run()
        assert not app.exception and not app.error
    assert len(service(remote).snapshot().matches) == 1


def test_two_sessions_preserve_both_saves():
    from concurrent.futures import ThreadPoolExecutor
    remote = RemoteSheet(); service(remote).snapshot()
    sessions = [service(remote), service(remote)]
    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(lambda s: match(s), sessions))
    snap = service(remote).snapshot()
    assert len(snap.matches) == 2 and len(set(snap.matches['ID'])) == 2
