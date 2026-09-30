from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from finanze_app import analytics as stats
from finanze_app.errors import ConflictError, StorageError, ValidationError
from finanze_app.service import FinanceService
from finanze_app.storage import SCHEMA
from finanze_app.utils import today


def movement(service, **changes):
    values = dict(kind="Uscita", amount=12.5, date=today(), description="Spesa", category="Cibo", source="Supermercato", account="Conto personale", notes="")
    values.update(changes)
    return service.save_movement(**values)


def match(service, **changes):
    values = dict(date=today(), package="0012", home="Squadra A", away="Squadra B", fee=45.0, account="Conto personale", category="Arbitraggio", notes="Rimborso")
    values.update(changes)
    return service.save_match(**values)


def test_initial_workbook_is_empty_and_typed(service):
    snapshot = service.snapshot()
    assert snapshot.movements.empty and snapshot.matches.empty
    assert stats.total_balance(snapshot) == 0
    workbook = load_workbook(service.store.path)
    assert workbook.sheetnames == list(SCHEMA)
    for name, columns in SCHEMA.items():
        assert [cell.value for cell in workbook[name][1]] == columns
        assert workbook[name].freeze_panes == "B2"
    assert workbook["Conti"]["C2"].value == 0
    assert workbook["Conti"]["C2"].data_type == "n"
    workbook.close()


def test_missing_sheets_are_created_without_overwriting_existing_data(service):
    movement(service)
    workbook = load_workbook(service.store.path)
    del workbook["Arbitraggio"]
    workbook.create_sheet("Personale")["A1"] = "Conservare"
    workbook.save(service.store.path)
    workbook.close()
    assert len(service.snapshot().movements) == 1
    workbook = load_workbook(service.store.path)
    assert "Arbitraggio" in workbook.sheetnames
    assert workbook["Personale"]["A1"].value == "Conservare"
    workbook.close()


def test_movement_crud_and_opening_balances(service):
    account = service.snapshot().accounts[0]
    service.save_account(name=account["Nome"], opening_balance=500, account_id=account["ID"])
    item_id = movement(service)
    assert stats.total_balance(service.snapshot()) == 487.5
    service.save_movement(kind="Entrata", amount=20, date=today(), description="Rimborso", category="Altro", source="Amico", account="Contanti", movement_id=item_id)
    assert stats.total_balance(service.snapshot()) == 520
    service.delete_movement(item_id)
    assert stats.total_balance(service.snapshot()) == 500


def test_match_creates_one_income_and_edit_preserves_link(service):
    match_id = match(service)
    snapshot = service.snapshot()
    first = snapshot.tables["Arbitraggio"][0]
    created = snapshot.tables["Movimenti"][0]["Data_creazione"]
    assert len(snapshot.movements) == len(snapshot.matches) == 1
    assert stats.total_balance(snapshot) == 45
    service.save_match(date=today(), package="0013", home="A", away="C", fee=60.5, account="Contanti", category="Altro", notes="Modificata", match_id=match_id, expected=first)
    snapshot = service.snapshot()
    entry = snapshot.tables["Movimenti"][0]
    assert len(snapshot.movements) == 1
    assert entry["ID"] == first["Movimento_ID"]
    assert entry["Arbitraggio_ID"] == match_id
    assert entry["Data_creazione"] == created
    assert entry["Importo"] == 60.5 and entry["Conto"] == "Contanti"
    assert entry["Note"] == "Modificata" and "Pacco 0013" in entry["Descrizione"]
    assert stats.total_balance(snapshot) == 60.5
    assert stats.match_totals(snapshot.matches) == (1, 60.5, 60.5)


@pytest.mark.parametrize("delete_from", ["match", "movement"])
def test_deleting_either_side_removes_the_match_and_income(service, delete_from):
    match_id = match(service)
    entry = service.snapshot().tables["Movimenti"][0]
    if delete_from == "match":
        service.delete_match(match_id)
    else:
        service.delete_movement(entry["ID"])
    snapshot = service.snapshot()
    assert snapshot.matches.empty and snapshot.movements.empty
    assert stats.total_balance(snapshot) == 0


def test_match_income_cannot_be_edited_independently(service):
    match(service)
    row = service.snapshot().tables["Movimenti"][0]
    before = service.store.path.read_bytes()
    with pytest.raises(ValidationError, match="Arbitraggio"):
        movement(service, movement_id=row["ID"])
    assert service.store.path.read_bytes() == before


def test_zero_fee_counts_as_a_match_but_not_as_income(service):
    match(service, fee=0)
    snapshot = service.snapshot()
    assert stats.match_totals(snapshot.matches) == (1, 0, 0)
    assert stats.totals(snapshot.movements) == (0, 0, 0)


@pytest.mark.parametrize("changes", [
    {"amount": -1}, {"amount": 0}, {"amount": float("nan")}, {"amount": 1.234},
    {"description": " "}, {"category": "Non esiste"}, {"account": "Non esiste"},
    {"date": today() + timedelta(days=1)}, {"kind": "Altro"},
])
def test_invalid_movement_never_changes_excel(service, changes):
    before = service.store.path.read_bytes()
    with pytest.raises(ValidationError):
        movement(service, **changes)
    assert service.store.path.read_bytes() == before


def test_invalid_match_does_not_write_either_sheet(service):
    before = service.store.path.read_bytes()
    with pytest.raises(ValidationError):
        match(service, home="")
    assert service.store.path.read_bytes() == before


def test_stale_form_is_rejected(service):
    item_id = movement(service)
    original = service.snapshot().tables["Movimenti"][0]
    movement(service, movement_id=item_id, amount=25)
    with pytest.raises(ConflictError):
        service.delete_movement(item_id, expected=original)
    assert service.snapshot().tables["Movimenti"][0]["Importo"] == 25


def test_category_rename_and_delete_update_both_sheets(service):
    match(service)
    original = next(row for row in service.snapshot().categories if row["Nome"] == "Arbitraggio")
    service.save_category(name="Partite", kind="Entrata", category_id=original["ID"], expected=original)
    snapshot = service.snapshot()
    assert snapshot.tables["Movimenti"][0]["Categoria"] == "Partite"
    assert snapshot.tables["Arbitraggio"][0]["Categoria"] == "Partite"
    service.delete_category(original["ID"], replacement="Altro")
    snapshot = service.snapshot()
    assert snapshot.tables["Movimenti"][0]["Categoria"] == "Altro"
    assert snapshot.tables["Arbitraggio"][0]["Categoria"] == "Altro"
    assert stats.total_balance(snapshot) == 45


def test_all_categories_can_be_deleted_without_losing_finances(service):
    match(service)
    movement(service)
    for row in service.snapshot().categories:
        service.delete_category(row["ID"])
    snapshot = service.snapshot()
    assert snapshot.categories == []
    assert set(snapshot.movements["Categoria"]) == {""}
    assert stats.total_balance(snapshot) == 32.5


def test_category_types_are_enforced_and_names_are_case_insensitive(service):
    movement(service)
    cibo = next(row for row in service.snapshot().categories if row["Nome"] == "Cibo")
    with pytest.raises(ValidationError):
        service.save_category(name="Cibo", kind="Entrata", category_id=cibo["ID"])
    with pytest.raises(ValidationError):
        service.save_category(name="cibo", kind="Entrambe")
    with pytest.raises(ValidationError):
        movement(service, category="Stipendio")
    category_id = service.save_category(name="Trasporti", kind="Entrambe")
    service.delete_category(category_id, replacement="Stipendio")


def test_account_rename_and_delete_preserve_total_and_match_links(service):
    account = service.snapshot().accounts[0]
    service.save_account(name="Principale", opening_balance=100, account_id=account["ID"])
    match(service, account="Principale")
    movement(service, account="Principale")
    original_total = stats.total_balance(service.snapshot())
    with pytest.raises(ValidationError):
        service.delete_account(account["ID"])
    service.delete_account(account["ID"], replacement="Contanti")
    snapshot = service.snapshot()
    assert snapshot.accounts[0]["Saldo_iniziale"] == 100
    assert set(snapshot.movements["Conto"]) == {"Contanti"}
    assert snapshot.tables["Arbitraggio"][0]["Conto"] == "Contanti"
    assert stats.total_balance(snapshot) == original_total == 132.5


def test_empty_accounts_can_all_be_deleted_and_recreated(service):
    for row in service.snapshot().accounts:
        service.delete_account(row["ID"])
    assert stats.total_balance(service.snapshot()) == 0
    with pytest.raises(ValidationError):
        movement(service)
    service.save_account(name="Nuovo", opening_balance=10)
    movement(service, account="Nuovo", amount=1)
    assert stats.total_balance(service.snapshot()) == 9


def test_exact_cents_and_monthly_periods(service):
    movement(service, amount=.1, date=date(2026, 1, 31))
    movement(service, amount=.2, date=date(2026, 2, 1))
    frame = service.snapshot().movements
    assert stats.totals(frame) == (0, .3, -.3)
    january = stats.date_filter(frame, date(2026, 1, 1), date(2026, 1, 31))
    assert stats.totals(january) == (0, .1, -.1)
    monthly = stats.monthly_flows(frame)
    assert monthly["Uscite"].tolist() == [.1, .2]


def test_period_chart_keeps_prior_balance(service):
    account = service.snapshot().accounts[0]
    service.save_account(name=account["Nome"], opening_balance=100, account_id=account["ID"])
    movement(service, kind="Entrata", category="Stipendio", amount=50, date=date(2026, 1, 1))
    movement(service, amount=10, date=date(2026, 2, 1))
    history = stats.balance_history(service.snapshot(), start=date(2026, 2, 1), end=date(2026, 2, 2))
    assert history["Saldo"].tolist() == [140, 140]


def test_backup_is_previous_valid_workbook(service):
    before = service.store.path.read_bytes()
    movement(service)
    assert service.store.backup_path.read_bytes() == before


def test_failed_save_leaves_original_file_and_no_partial_transaction(service, monkeypatch):
    import finanze_app.storage as module
    before = service.store.path.read_bytes()
    original_replace = module.os.replace
    def fail_target(source, target):
        if Path(target) == service.store.path:
            raise PermissionError("Locked by Excel")
        return original_replace(source, target)
    monkeypatch.setattr(module.os, "replace", fail_target)
    with pytest.raises(StorageError):
        match(service)
    assert service.store.path.read_bytes() == before
    assert service.snapshot().matches.empty


def test_parallel_sessions_do_not_lose_rows(service):
    path = service.store.path
    def write(index):
        return movement(FinanceService(path), description=f"Spesa {index}", amount=1)
    with ThreadPoolExecutor(max_workers=4) as pool:
        ids = list(pool.map(write, range(12)))
    snapshot = service.snapshot()
    assert len(set(ids)) == len(snapshot.movements) == 12
    assert stats.total_balance(snapshot) == -12


def test_formula_like_notes_are_literal_excel_text(service):
    movement(service, notes='=HYPERLINK("https://example.invalid", "test")')
    workbook = load_workbook(service.store.path)
    assert workbook["Movimenti"]["H2"].data_type == "s"
    workbook.close()


def test_bad_excel_is_not_reset_or_overwritten(service):
    service.store.path.write_bytes(b"not an xlsx")
    with pytest.raises(StorageError):
        service.snapshot()
    assert service.store.path.read_bytes() == b"not an xlsx"
