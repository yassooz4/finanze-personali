from datetime import timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from finanze_app.analytics import total_balance
from finanze_app.errors import ConflictError, ValidationError
from finanze_app.utils import today


def add(service, **changes):
    values = dict(date=today() - timedelta(days=40), package="0008", home="Casa", away="Ospiti", fee=45, category="Arbitraggio")
    values.update(changes)
    return service.save_match(**values)


def payment(service, record, status="Ricevuto", **changes):
    values = {"ID": record["ID"], "Stato": status, "Data incasso": today(), "Conto": "Conto personale", "expected": record}
    values.update(changes)
    service.update_match_payments([values])


def test_new_match_is_pending_and_never_changes_balance(service):
    add(service)
    snapshot = service.snapshot()
    assert total_balance(snapshot) == 0
    assert snapshot.movements.empty
    record = snapshot.tables["Arbitraggio"][0]
    assert record["Stato"] == "Da ricevere"
    assert record["Data incasso"] is None and record["Movimento_ID"] == ""
    assert len(snapshot.matches) == 1


def test_receipt_uses_payment_date_repeated_save_no_duplicates_and_reversal(service):
    add(service)
    record = service.snapshot().tables["Arbitraggio"][0]
    payment(service, record)
    snapshot = service.snapshot()
    paid = snapshot.tables["Arbitraggio"][0]
    movement = snapshot.tables["Movimenti"][0]
    assert total_balance(snapshot) == 45 and movement["Data"] == today()
    assert movement["Data"] != paid["Data partita"]
    payment(service, paid)
    assert len(service.snapshot().movements) == 1
    assert service.snapshot().tables["Movimenti"][0]["ID"] == movement["ID"]
    paid = service.snapshot().tables["Arbitraggio"][0]
    payment(service, paid, "Da ricevere")
    snapshot = service.snapshot()
    assert total_balance(snapshot) == 0 and snapshot.movements.empty
    assert snapshot.tables["Arbitraggio"][0]["Data incasso"] is None
    payment(service, snapshot.tables["Arbitraggio"][0])
    assert len(service.snapshot().movements) == 1 and total_balance(service.snapshot()) == 45


@pytest.mark.parametrize("changes", [{"Data incasso": None}, {"Data incasso": today()+timedelta(days=1)}, {"Data incasso": today()-timedelta(days=41)}, {"Conto": ""}, {"Stato": "invalid"}])
def test_invalid_receipt_is_atomic(service, changes):
    add(service)
    record = service.snapshot().tables["Arbitraggio"][0]
    before = service.store.path.read_bytes()
    with pytest.raises(ValidationError):
        payment(service, record, **changes)
    assert service.store.path.read_bytes() == before
    assert service.snapshot().movements.empty


def test_batch_payments_all_or_nothing_and_stale_state(service):
    add(service)
    add(service, package="0009")
    a,b = service.snapshot().tables["Arbitraggio"]
    updates = [{"ID": r["ID"], "Stato": "Ricevuto", "Data incasso": today(), "Conto": "Conto personale", "expected": r} for r in [a,b]]
    updates[1]["Conto"] = "missing"
    with pytest.raises(ValidationError):
        service.update_match_payments(updates)
    assert service.snapshot().movements.empty
    payment(service,a)
    with pytest.raises(ConflictError):
        payment(service,a)
    assert len(service.snapshot().movements) == 1


def test_legacy_excel_migration_preserves_ids_balance_and_dates(service):
    add(service,status="Ricevuto",received_date=today(),account="Contanti")
    before = service.snapshot()
    entry = before.tables["Movimenti"][0]
    book = load_workbook(service.store.path)
    sheet = book["Arbitraggio"]
    headers = [c.value for c in sheet[1]]
    for name in ["Data incasso", "Stato"]:
        sheet.delete_cols(headers.index(name)+1)
        headers = [c.value for c in sheet[1]]
    book.save(service.store.path)
    book.close()
    snapshot = service.snapshot()
    record = snapshot.tables["Arbitraggio"][0]
    assert record["Stato"] == "Ricevuto" and record["Data incasso"] == entry["Data"]
    assert record["Movimento_ID"] == entry["ID"] and total_balance(snapshot) == 45
    payment(service,record)
    again = service.snapshot()
    assert total_balance(again) == 45 and len(again.movements) == 1
    assert again.tables["Arbitraggio"][0]["Stato"] == "Ricevuto"


def test_delete_pending_match_and_pending_without_accounts(service):
    for account in service.snapshot().accounts:
        service.delete_account(account["ID"])
    mid = add(service)
    assert len(service.snapshot().matches)==1
    service.delete_match(mid)
    assert service.snapshot().matches.empty and service.snapshot().movements.empty


def test_ui_defaults_pending_and_table_receipt(service, monkeypatch):
    monkeypatch.setenv("FINANZE_FILE", str(service.store.path))
    at = AppTest.from_file(str(Path(__file__).resolve().parents[1]/"app.py"), default_timeout=30)
    at.secrets["accesso"] = {"password":"test"}
    at.run()
    at.text_input(key="_login_password").set_value("test")
    at.button[0].click().run()
    at.sidebar.radio[0].set_value("⚽ Arbitraggio").run()
    at.button(key="top_new").click().run()
    def widget(elements,label):
        return next(e for e in elements if e.label==label)
    assert widget(at.selectbox,"Stato pagamento").value=="Da ricevere"
    widget(at.text_input,"Numero pacco").set_value("09")
    widget(at.text_input,"Squadra di casa").set_value("A")
    widget(at.text_input,"Squadra ospite").set_value("B")
    widget(at.number_input,"Compenso previsto (€)").set_value(50)
    widget(at.number_input,"Km percorsi").set_value(32.5)
    widget(at.text_input,"Categoria della partita").set_value("Juniores")
    widget(at.button,"Salva partita").click().run()
    assert not at.exception and service.snapshot().movements.empty
    # Simulate the table's editable-cell event, then submit the actual save action.
    keys = [k for k in at.session_state._state.filtered_state if k.startswith("referee_payments_")]
    assert keys
    at.session_state[keys[-1]] = {"edited_rows": {0: {"Stato":"Ricevuto","Data incasso":today().isoformat(),"Conto":"Contanti"}},"added_rows":[],"deleted_rows":[]}
    at.button(key="save_payments").click().run()
    assert not at.exception
    assert total_balance(service.snapshot())==50
    record = service.snapshot().tables["Arbitraggio"][0]
    assert record["Stato"]=="Ricevuto"
    assert record["Km"]==32.5 and record["Categoria partita"]=="Juniores"
    assert record["Categoria"]=="Arbitraggio"



def test_match_details_edit_and_excel_migration_preserve_receipt(service):
    mid = add(service, km=32.5, match_category="Allievi", status="Ricevuto", received_date=today(), account="Contanti")
    original = service.snapshot().tables["Arbitraggio"][0]
    entry_id = original["Movimento_ID"]
    add(service, match_id=mid, km=40, match_category="Juniores", status="Ricevuto", received_date=today(), account="Contanti")
    record = service.snapshot().tables["Arbitraggio"][0]
    assert record["Km"] == 40 and record["Categoria partita"] == "Juniores"
    assert record["Categoria"] == "Arbitraggio" and record["Movimento_ID"] == entry_id
    assert total_balance(service.snapshot()) == 45 and len(service.snapshot().movements) == 1
    # Omitted metadata stays intact when a caller edits other match fields.
    add(service, match_id=mid, status="Ricevuto", received_date=today(), account="Contanti")
    record = service.snapshot().tables["Arbitraggio"][0]
    assert record["Km"] == 40 and record["Categoria partita"] == "Juniores"
    workbook = load_workbook(service.store.path)
    sheet = workbook["Arbitraggio"]
    headers = [c.value for c in sheet[1]]
    assert sheet.cell(2, headers.index("Km")+1).data_type == "n"
    for name in ["Categoria partita", "Km"]:
        headers = [c.value for c in sheet[1]]
        sheet.delete_cols(headers.index(name)+1)
    workbook.save(service.store.path)
    workbook.close()
    migrated = service.snapshot()
    assert migrated.tables["Arbitraggio"][0]["Km"] == 0
    assert migrated.tables["Arbitraggio"][0]["Categoria partita"] == ""
    assert total_balance(migrated) == 45
    assert migrated.tables["Movimenti"][0]["ID"] == entry_id


def test_negative_km_is_rejected_atomically(service):
    before = service.store.path.read_bytes()
    with pytest.raises(ValidationError):
        add(service, km=-1)
    assert service.store.path.read_bytes() == before
