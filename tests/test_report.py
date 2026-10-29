from datetime import date
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from finanze_app import report as reports


def seed(service):
    personal = next(row for row in service.snapshot().accounts if row["Nome"] == "Conto personale")
    service.save_account(name="Conto personale", opening_balance=100, account_id=personal["ID"], expected=personal)
    common = dict(source="Test", account="Conto personale")
    service.save_movement(kind="Entrata", amount=1000, date=date(2026, 1, 10), description="Stipendio gennaio", category="Stipendio", **common)
    service.save_movement(kind="Uscita", amount=50.5, date=date(2026, 1, 15), description="Spesa", category="Cibo", **common)
    service.save_movement(kind="Uscita", amount=20, date=date(2026, 2, 3), description="Pizza", category="Cibo", **common)
    service.save_movement(kind="Giroconto", amount=30, date=date(2026, 2, 5), description="", category="", source="", account="Conto personale", destination_account="Contanti")
    return service.snapshot()


def summary(report):
    return {label: value for label, value, _ in report.summary}


def test_full_report_totals_balances_and_running_balance(service):
    report = reports.build_report(seed(service))
    values = summary(report)
    assert values["Entrate"] == 1000 and values["Uscite"] == 70.5
    assert values["Saldo a inizio periodo"] == 100
    assert values["Saldo a fine periodo"] == values["Saldo attuale"] == 1029.5
    assert values["Movimenti nel report"] == 4
    assert report.movements["Saldo progressivo"].tolist() == [1100, 1049.5, 1029.5, 1029.5]
    assert report.movements["Importo"].tolist() == [1000, -50.5, -20, 30]
    assert set(report.accounts["Conto"]) >= {"Conto personale", "Contanti"}


def test_month_filter_uses_balance_at_period_edges(service):
    snapshot = seed(service)
    report = reports.build_report(snapshot, start=date(2026, 2, 1), end=date(2026, 2, 28))
    values = summary(report)
    assert values["Saldo a inizio periodo"] == 1049.5
    assert values["Saldo a fine periodo"] == 1029.5
    assert values["Uscite"] == 20 and values["Entrate"] == 0
    assert len(report.movements) == 2
    assert report.months["Mese"].tolist() == ["Febbraio 2026"]
    row = report.accounts.set_index("Conto").loc["Conto personale"]
    assert row["Saldo iniziale periodo"] == 1049.5 and row["Saldo finale periodo"] == 999.5  # 30 € passati a Contanti con il giroconto


def test_account_view_shows_transfers_as_cash_flow(service):
    report = reports.build_report(seed(service), account="Contanti")
    assert report.movements["Tipo"].tolist() == ["Giroconto in entrata"]
    values = summary(report)
    assert values["Entrate"] == 30 and values["Saldo attuale"] == 30
    assert report.categories.empty


def test_kind_and_category_filters_do_not_change_balances(service):
    snapshot = seed(service)
    report = reports.build_report(snapshot, kind="Uscita", category="Cibo")
    values = summary(report)
    assert len(report.movements) == 2 and values["Uscite"] == 70.5 and values["Entrate"] == 0
    assert values["Saldo attuale"] == 1029.5
    assert report.categories.to_dict("records") == [{"Tipo": "Uscita", "Categoria": "Cibo", "Importo": 70.5, "% sul totale": 1.0, "Movimenti": 2}]
    assert len(reports.build_report(snapshot, kind="Giroconto").movements) == 1


def test_empty_period_still_has_balances(service):
    report = reports.build_report(seed(service), start=date(2026, 6, 1), end=date(2026, 6, 30))
    assert report.movements.empty and report.categories.empty and report.months.empty
    values = summary(report)
    assert values["Saldo a inizio periodo"] == values["Saldo a fine periodo"] == 1029.5


def test_excel_and_csv_outputs(service):
    report = reports.build_report(seed(service))
    workbook = load_workbook(BytesIO(reports.to_excel(report)))
    assert workbook.sheetnames == reports.SECTIONS
    assert workbook["Movimenti"].max_row == 5
    assert workbook["Movimenti"]["A2"].value.date() == date(2026, 1, 10)
    only = load_workbook(BytesIO(reports.to_excel(report, ["Riepilogo", "Mesi"])))
    assert only.sheetnames == ["Riepilogo", "Mesi"]
    lines = reports.to_csv(report).decode("utf-8-sig").splitlines()
    assert lines[0].startswith("Data;Tipo;Importo") and len(lines) == 5


def test_text_starting_with_equals_is_never_a_formula():
    frame = pd.DataFrame({"Descrizione": ["=1+1"], "Importo": [5.0]})
    report = reports.Report([], frame, frame.iloc[0:0], frame.iloc[0:0], frame.iloc[0:0], frame.iloc[0:0])
    cell = load_workbook(BytesIO(reports.to_excel(report, ["Movimenti"])))["Movimenti"]["A2"]
    assert cell.data_type == "s" and cell.value == "=1+1"
